from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from scripts.columns import (
    ACCESSION,
    CAS,
    CID,
    COLUMNS_BY_GROUP,
    DOI,
    GENE_NAME,
    MIXTURE,
    MOLECULE,
    MOLECULE_NAME,
    PARAMETER,
    RECEPTOR,
    RECEPTOR_NAME,
    RESPONSE,
    SEQUENCE,
    SOURCE,
    SPECIES,
    UNIPROT_ID,
)
from scripts.fetch_data import _BlastCandidate, _select_blast_candidates_non_interactive
from scripts.process_molecules import validate_cid_or_cas
from scripts.process_receptors import (
    _review_receptor_names,
    _review_species_by_accession,
    process_receptors_name_columns,
    suggest_canonical_receptor_name,
)
from scripts.process_responses import validate_parameter_column
from scripts.process_sources import normalize_doi_column
from scripts.read_excel import get_raw_data_from_excel_file
from scripts.report import (
    ExcelLayout,
    Issue,
    IssueCollector,
    ValidationError,
    build_report,
    emit,
    excel_row,
    merge_issues,
)


def _frame(columns: dict[tuple[str, str], list[object]]) -> pd.DataFrame:
    return pd.DataFrame({k: v for k, v in columns.items()})


def _collect(step, *args) -> IssueCollector:  # noqa: ANN001
    collector = IssueCollector(catch_unexpected=True)
    with collector.stage("test"):
        step(*args)
    return collector


# --- report core -------------------------------------------------------------


def test_excel_row_accounts_for_two_header_rows():
    assert excel_row(0) == 3
    assert excel_row(np.int64(10)) == 13


def test_emit_outside_a_collector_is_a_no_op():
    emit(Issue(code="x", message="ignored"))


def test_stage_records_validation_error_and_continues():
    collector = IssueCollector(catch_unexpected=False)
    with collector.stage("a"):
        raise ValidationError("boom", [Issue(code="bad", message="m", rows=[1])])
    with collector.stage("b"):
        emit(Issue(code="note", message="n", severity="warning"))
    assert collector.stages == {"a": "failed", "b": "passed"}
    assert collector.failure_messages == ["boom"]
    assert [i.stage for i in collector.issues] == ["a", "b"]
    assert collector.has_errors


def test_stage_wraps_plain_value_error():
    collector = IssueCollector(catch_unexpected=False)
    with collector.stage("a"):
        raise ValueError("legacy")
    assert collector.issues[0].code == "unstructured_error"
    assert collector.failure_messages == ["legacy"]


def test_unexpected_error_recorded_only_when_asked():
    collector = IssueCollector(catch_unexpected=True)
    with collector.stage("a"):
        raise KeyError("x")
    assert collector.crashed
    assert collector.issues[0].code == "internal_error"

    strict = IssueCollector(catch_unexpected=False)
    with pytest.raises(KeyError), strict.stage("a"):
        raise KeyError("x")


def test_check_keeps_going_after_a_failed_step():
    collector = IssueCollector(catch_unexpected=False)

    def failing() -> None:
        raise ValidationError("first", [Issue(code="one", message="m")])

    with collector.stage("s"):
        assert not collector.check(failing)
        assert collector.check(lambda: None)
    assert collector.stages["s"] == "failed"


def test_merge_issues_unions_rows_of_identical_issues():
    merged = merge_issues(
        [
            Issue(code="c", message="m", rows=[4], value="v"),
            Issue(code="c", message="m", rows=[1], value="v"),
            Issue(code="c", message="m", rows=[2], value="other"),
        ]
    )
    assert [(i.value, i.rows) for i in merged] == [("v", [1, 4]), ("other", [2])]


def test_layout_maps_renamed_columns_back_to_excel_headers():
    raw = _frame(
        {
            (RECEPTOR, "Order"): ["Diptera", "Diptera"],
            (RECEPTOR, GENE_NAME): ["OR5", "Or6"],
            (RECEPTOR, UNIPROT_ID): ["P1", "P2"],
        }
    )
    raw.attrs["sheet_name"] = "Feuil1"
    layout = ExcelLayout.from_raw(raw)
    assert layout.excel_column(RECEPTOR, RECEPTOR_NAME) == "B"
    assert layout.excel_column(RECEPTOR, ACCESSION) == "C"
    assert layout.excel_column(RECEPTOR, "Sequence_ref") is None
    assert layout.raw_values(RECEPTOR, RECEPTOR_NAME, [0, 1]) == ["OR5", "Or6"]

    collector = IssueCollector(catch_unexpected=False)
    with collector.stage("receptors"):
        emit(
            Issue(
                code="receptor_name_format",
                message="m",
                severity="auto_fix",
                group=RECEPTOR,
                column=RECEPTOR_NAME,
                rows=[0],
                value=np.str_("OR5"),
                suggested_value="Or5",
                details={"identity": np.float64("nan")},
            )
        )
    report = build_report(
        collector,
        mode="check",
        study_id="s",
        excel_path="s.xlsx",
        file_sha256="0",
        layout=layout,
        blast_mode="cached",
    )
    assert report["status"] == "passed"
    issue = report["issues"][0]
    assert issue["column"] == GENE_NAME
    assert issue["cells"] == ["B3"]
    assert issue["sheet"] == "Feuil1"
    assert issue["value"] == "OR5"
    assert issue["details"] == {"identity": None}


# --- structure ---------------------------------------------------------------


def _write_template(path: Path, rename: dict[str, str] | None = None) -> None:
    rename = rename or {}
    wb = openpyxl.Workbook()
    wb.active.title = "about"
    ws = wb.create_sheet("data")
    col = 1
    for group, columns in COLUMNS_BY_GROUP:
        for i, name in enumerate(columns):
            ws.cell(row=1, column=col, value=group if i == 0 else None)
            ws.cell(row=2, column=col, value=rename.get(name, name))
            col += 1
    ws.cell(row=3, column=1, value="Diptera")
    wb.save(path)


def test_read_excel_keeps_sheet_name(tmp_path: Path):
    path = tmp_path / "ok.xlsx"
    _write_template(path)
    df = get_raw_data_from_excel_file(path)
    assert df.attrs["sheet_name"] == "data"
    assert df.loc[0, (RECEPTOR, "Order")] == "Diptera"


def test_missing_column_points_at_misspelled_header(tmp_path: Path):
    path = tmp_path / "bad.xlsx"
    _write_template(path, {"Solvent used for dilution": "Solvant used for dilution"})
    with pytest.raises(ValidationError) as excinfo:
        get_raw_data_from_excel_file(path)
    (issue,) = excinfo.value.issues
    assert issue.code == "missing_column"
    assert issue.value == "Solvant used for dilution"
    assert issue.suggested_value == "Solvent used for dilution"
    # Receptor (7) + Co-Receptor (6) + 6th Molecule column -> column S
    assert issue.cells == ["S2"]


# --- per-column validations --------------------------------------------------


def test_value_not_allowed_lists_rows_and_closest_value():
    df = _frame({(RESPONSE, PARAMETER): ["primary", "primery", None, "primery", "ec50"]})
    with pytest.raises(ValidationError) as excinfo:
        validate_parameter_column(df)
    issues = {i.value: i for i in excinfo.value.issues}
    assert issues["primery"].rows == [1, 3]
    assert issues["primery"].suggested_value == "primary"
    assert issues["ec50"].suggested_value is None


def test_cid_and_cas_missing_issue_per_row():
    df = _frame(
        {
            (MOLECULE, MOLECULE_NAME): ["a", "b", "c"],
            (MOLECULE, CID): ["123", None, "mixture of stuff"],
            (MOLECULE, CAS): [None, None, None],
            (MOLECULE, MIXTURE): [None, None, "mixture"],
        }
    )
    with pytest.raises(ValidationError) as excinfo:
        validate_cid_or_cas(df)
    assert [(i.rows, i.column) for i in excinfo.value.issues] == [([1], CID)]


def test_doi_url_prefix_is_reported_as_auto_fix():
    df = _frame({(SOURCE, DOI): ["https://doi.org/10.1/abc", "10.1/abc"]})
    collector = _collect(normalize_doi_column, df)
    (issue,) = collector.issues
    assert (issue.code, issue.rows, issue.suggested_value) == (
        "doi_url_prefix",
        [0],
        "10.1/abc",
    )


# --- receptor names ----------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "expected"),
    [("BimOR34", "Or34"), ("AgamOrco", "Orco"), ("Or22A", "Or22a"), ("XP_1", None)],
)
def test_suggest_canonical_receptor_name(name: str, expected: str | None):
    assert suggest_canonical_receptor_name(name) == expected


def _review(original: list[object], accessions: list[object], sequences: list[object],
            excel_accessions: list[object] | None = None):  # noqa: ANN202
    df = _frame(
        {
            (RECEPTOR, RECEPTOR_NAME): [None] * len(original),
            (RECEPTOR, ACCESSION): accessions,
            (RECEPTOR, SEQUENCE): sequences,
        }
    )
    collector = IssueCollector(catch_unexpected=False)
    with collector.stage("receptors"):
        blocking = _review_receptor_names(
            df,
            [RECEPTOR],
            {RECEPTOR: pd.Series(original)},
            {RECEPTOR: pd.Series(excel_accessions)} if excel_accessions else None,
        )
    return blocking, collector.issues


def test_receptor_names_come_from_the_excel_only():
    blocking, emitted = _review(
        ["OR5", "BimOR34", "Or125", None, "Or93", None],
        ["P5", "P34", "KM1", "P7", "KM1", None],
        ["MA", "MB", "MC", "MD", "MC", None],
    )
    by_row = {i.rows[0]: i for i in blocking}
    (fmt,) = emitted
    assert (fmt.rows, fmt.code, fmt.suggested_value) == ([0], "receptor_name_format", "Or5")
    assert by_row[1].code == "receptor_name_invalid"
    assert by_row[1].suggested_value == "Or34"  # from the Excel text itself
    assert by_row[3].code == "receptor_name_missing"
    for row in (2, 4):
        assert by_row[row].code == "receptor_accession_shared"
        assert by_row[row].details["names_for_accession"] == ["Or125", "Or93"]
        assert by_row[row].details["identical_sequences"] is True
    assert 5 not in by_row  # no receptor on that row


def test_same_name_in_different_case_is_not_an_accession_conflict():
    blocking, emitted = _review(["OR93", "Or93"], ["KM1", "KM1"], ["MC", "MC"])
    assert blocking == []
    assert [i.code for i in emitted] == ["receptor_name_format"]


def test_accession_shared_through_blast_is_only_a_warning():
    blocking, emitted = _review(
        ["Or182", "Or183"], ["XP_1", "XP_1"], ["MDIL", "MDIV"], [None, None]
    )
    assert blocking == []
    assert {i.code for i in emitted} == {"receptor_blast_hit_shared"}
    assert emitted[0].details["identical_sequences"] is False


def test_names_are_never_taken_from_the_reference_record():
    df = _frame(
        {
            (RECEPTOR, RECEPTOR_NAME): ["OR5", "BimOR34", "Or1", "Or2"],
            (RECEPTOR, ACCESSION): ["P5", "P34", "Not available", "Not available"],
            (RECEPTOR, SEQUENCE): ["MA", "MB", "MC", "MD"],
        }
    )
    with pytest.raises(ValidationError) as excinfo:
        process_receptors_name_columns(
            df, [RECEPTOR], ["Not available", "P34", "P5"]
        )
    assert [i.code for i in excinfo.value.issues] == ["receptor_name_invalid"]
    # The placeholder accession is shared by unrelated receptors: each keeps
    # its own name.
    names = df[RECEPTOR][RECEPTOR_NAME]
    assert pd.isna(names[1])
    assert names.drop(1).tolist() == ["Or5", "Or1", "Or2"]


def test_an_accession_given_several_species_is_blocking():
    df = _frame(
        {
            (RECEPTOR, ACCESSION): [
                "ADB89179.1",
                "ADB89179.1",
                "P1",
                "Not available",
                "Not available",
            ],
            (RECEPTOR, SPECIES): [
                "Ostrinia nubilalis",
                "Ostrinia furnacalis",
                "Ostrinia nubilalis",
                "A",
                "B",
            ],
        }
    )
    issues = _review_species_by_accession(df, [RECEPTOR])
    assert [(i.code, i.value, i.rows) for i in issues] == [
        ("receptor_accession_species_conflict", "Ostrinia nubilalis", [0]),
        ("receptor_accession_species_conflict", "Ostrinia furnacalis", [1]),
    ]
    assert issues[0].details["species_for_accession"] == [
        "Ostrinia furnacalis",
        "Ostrinia nubilalis",
    ]


# --- BLAST -------------------------------------------------------------------


def test_cached_blast_mode_runs_nothing_and_reports_pending():
    candidates = [
        _BlastCandidate("new", "MKV", "Aedes aegypti", "Or1"),
        _BlastCandidate("no-hit", "MKW", "Aedes aegypti", "Or2", "no hit found"),
    ]
    collector = IssueCollector(catch_unexpected=False)
    with collector.stage("receptors"):
        assert _select_blast_candidates_non_interactive(candidates, "cached") == []
    (issue,) = collector.issues
    assert issue.code == "blast_pending"
    assert len(issue.details["queries"]) == 1


def test_run_blast_mode_skips_definitive_no_hits():
    candidates = [
        _BlastCandidate("new", "MKV", "sp", "Or1"),
        _BlastCandidate("retry", "MKX", "sp", "Or3", "timeout"),
        _BlastCandidate("no-hit", "MKW", "sp", "Or2", "no hit found"),
    ]
    selected = _select_blast_candidates_non_interactive(candidates, "run")
    assert [c.category for c in selected] == ["new", "retry"]
