"""Split an Excel input file into one file per unique DOI value.

Header rows (group labels + column names) are preserved in every output file.
Rows sharing the same DOI are written in their original order.
Output filenames follow the pattern: FirstAuthor_Journal-Year.xlsx,
derived from the "Reference" column ("Author(s) YEAR Journal").

Usage:
    python -m scripts.tools.split_by_doi --input path/to/file.xlsx
    python -m scripts.tools.split_by_doi --input path/to/file.xlsx --output-dir out/
"""

import argparse
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import TYPE_CHECKING, Any

import openpyxl

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

_HEADER_ROW_COUNT = 2
_DOI_COLUMN_NAME = "DOI"
_REFERENCE_COLUMN_NAME = "Reference"

# Matches: "FirstAuthor [et al / & X] YEAR Journal name"
_REFERENCE_RE = re.compile(r"^(\S+)\s+.*?(\d{4})\s+(.+)$")


def _find_column(ws: Worksheet, name: str) -> int:
    for cell in ws[_HEADER_ROW_COUNT]:
        if cell.value == name:
            return int(cell.column)  # type: ignore[arg-type]
    msg = f"Column '{name}' not found in row {_HEADER_ROW_COUNT}"
    raise ValueError(msg)


def _filename_from_reference(reference: str | None) -> str:
    """Parse 'Author(s) YEAR Journal' → 'Author_Journal-YEAR'."""
    if not reference:
        return "no_reference"
    m = _REFERENCE_RE.match(reference.strip())
    if not m:
        return re.sub(r'[<>:"/\\|?*\s]', "_", reference)[:120]
    author, year, journal = m.group(1), m.group(2), m.group(3)
    journal_clean = re.sub(r"\s+", "", journal)
    raw = f"{author}_{journal_clean}-{year}"
    return re.sub(r'[<>:"/\\|?*]', "_", raw)[:120]


def _copy_sheet_verbatim(src: Worksheet, dst: Worksheet) -> None:
    for r_idx, row in enumerate(src.iter_rows(), start=1):
        for c_idx, cell in enumerate(row, start=1):
            dst.cell(row=r_idx, column=c_idx, value=cell.value)
    for merge in src.merged_cells.ranges:
        dst.merge_cells(str(merge))


def split_by_doi(input_path: Path, output_dir: Path) -> None:
    wb = openpyxl.load_workbook(input_path)

    if len(wb.sheetnames) < 2:
        msg = f"Expected at least 2 sheets in {input_path}, found: {wb.sheetnames}"
        raise ValueError(msg)

    meta_sheet_name = wb.sheetnames[0]
    data_sheet_name = wb.sheetnames[1]
    ws_meta: Worksheet = wb[meta_sheet_name]  # type: ignore[assignment]
    ws_data: Worksheet = wb[data_sheet_name]  # type: ignore[assignment]

    doi_col = _find_column(ws_data, _DOI_COLUMN_NAME)
    ref_col = _find_column(ws_data, _REFERENCE_COLUMN_NAME)

    # Read the two header rows as lists of (column_index, value) pairs
    header_rows: list[list[tuple[int, Any]]] = [
        [(c_idx, cell.value) for c_idx, cell in enumerate(ws_data[row_num], start=1)]
        for row_num in range(1, _HEADER_ROW_COUNT + 1)
    ]

    # Collect merged regions that belong entirely to the header rows
    header_merges = [
        str(m) for m in ws_data.merged_cells.ranges if m.max_row <= _HEADER_ROW_COUNT
    ]

    # Group data rows (row 3+) by DOI, preserving insertion order
    # Also capture the first Reference value seen for each DOI
    groups: dict[str, list[list[tuple[int, Any]]]] = defaultdict(list)
    references: dict[str, str | None] = {}

    for row in ws_data.iter_rows(min_row=_HEADER_ROW_COUNT + 1):
        doi_val = row[doi_col - 1].value
        doi_key = str(doi_val).strip() if doi_val is not None else ""
        if doi_key not in references:
            ref_val = row[ref_col - 1].value
            references[doi_key] = str(ref_val).strip() if ref_val is not None else None
        row_data = [(c_idx, cell.value) for c_idx, cell in enumerate(row, start=1)]
        groups[doi_key].append(row_data)

    output_dir.mkdir(parents=True, exist_ok=True)

    used_names: dict[str, int] = {}

    for doi, data_rows in groups.items():
        out_wb = openpyxl.Workbook()

        out_ws_meta: Worksheet = out_wb.active  # type: ignore[assignment]
        out_ws_meta.title = meta_sheet_name
        _copy_sheet_verbatim(ws_meta, out_ws_meta)

        out_ws_data: Worksheet = out_wb.create_sheet(title=data_sheet_name)  # type: ignore[assignment]

        for r_idx, row in enumerate(header_rows, start=1):
            for c_idx, value in row:
                out_ws_data.cell(row=r_idx, column=c_idx, value=value)

        for merge_range in header_merges:
            out_ws_data.merge_cells(merge_range)

        for r_idx, row in enumerate(data_rows, start=_HEADER_ROW_COUNT + 1):
            for c_idx, value in row:
                out_ws_data.cell(row=r_idx, column=c_idx, value=value)

        base_name = _filename_from_reference(references.get(doi))
        count = used_names.get(base_name, 0)
        used_names[base_name] = count + 1
        out_name = f"{base_name}_{count + 1}" if count > 0 else base_name
        out_path = output_dir / f"{out_name}.xlsx"
        out_wb.save(out_path)
        print(f"  {len(data_rows):>4} rows  →  {out_path}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Split an Excel data file into one file per unique DOI."
    )
    parser.add_argument("--input", required=True, type=Path, help="Input Excel file")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=None,
        help="Output directory (default: <input_dir>/split_by_doi/)",
    )
    args = parser.parse_args()

    if not args.input.exists():
        print(f"Error: input file not found: {args.input}", file=sys.stderr)
        sys.exit(1)

    output_dir: Path = args.output_dir or (args.input.parent / "split_by_doi")

    print(f"Splitting {args.input} by DOI → {output_dir}/")
    split_by_doi(args.input, output_dir)
    print("Done.")


if __name__ == "__main__":
    main()
